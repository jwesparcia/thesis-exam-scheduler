import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from collections import defaultdict
import os
import re

def clean_str(val):
    if pd.isna(val) or val is None:
        return ""
    s = str(val).strip()
    if s.lower() == "nan":
        return ""
    # Fix common encoding artifacts
    s = s.replace("Pea", "Peña").replace("Pion", "Piñon")
    return s

def classify_subject(name: str):
    name_lower = name.lower()
    
    # Practical keywords
    practical_keywords = [
        "physical education", "national service training program", "euthenics", 
        "thesis", "practicum", "nstp", "immersion", "capstone", "laboratory",
        "pathfit", "p.e."
    ]
    if any(k in name_lower for k in practical_keywords):
        exam_type = "practical"
    else:
        exam_type = "written"

    # General education keywords
    general_keywords = [
        "oral communication", "general mathematics", "21st century literature",
        "reading and writing", "statistics and probability", "understanding the self",
        "understanding self", "contemporary world", "purposive communication", "ethics", 
        "art appreciation", "komunikasyon at pananaliksik", "pagbasa at pagsusuri", 
        "personal development", "philosophy", "literature", "media and information literacy",
        "euthenics", "national service training program", "rotc", "p.e./pathfit",
        "physical education", "readings in philippine history", "rizal", "philippine popular culture",
        "the entrepreneurial mind", "mathematics in the modern world", "science, technology",
        "great books", "foreign language", "general physics", "general chemistry", "general biology",
        "understanding culture, society", "intro to philosophy", "personal development",
        "disaster readiness", "earth science", "physical science", "earth and life science"
    ]
    if name.startswith("GE") or any(k in name_lower for k in general_keywords):
        category = "general"
    else:
        category = "major"

    return exam_type, category

def generate_aligned_curriculum():
    input_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../data/import/school_curriculum_sample_edited.xlsx"))
    output_path_1 = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../data/import/school_curriculum.xlsx"))
    output_path_2 = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../data/import/school_curriculum_sample_edited.xlsx"))
    
    xl = pd.ExcelFile(input_path)
    
    wb = openpyxl.Workbook()
    # remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Navy Blue
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Calibri", size=10)
    data_align_left = Alignment(horizontal="left", vertical="center")
    data_align_center = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    columns = [
        "Course", "Category", "Year Level", "Section", 
        "Subject Code", "Subject Name", "Teacher Name", 
        "Subject Category", "Exam Type"
    ]

    for sheet_name in ["1st Semester", "2nd Semester", "3rd Semester"]:
        ws = wb.create_sheet(title=sheet_name)
        
        # Write headers
        ws.row_dimensions[1].height = 28
        for col_idx, col_name in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        if sheet_name not in xl.sheet_names:
            continue

        df = pd.read_excel(xl, sheet_name=sheet_name)
        if df.empty:
            continue

        row_idx = 2
        for _, row in df.iterrows():
            sub_name = clean_str(row.get("Subject Name"))
            if not sub_name:
                continue
            course_val = clean_str(row.get("Course"))
            year_val = clean_str(row.get("Year Level"))
            section_val = clean_str(row.get("Section"))
            code_val = clean_str(row.get("Subject Code"))
            teacher_val = clean_str(row.get("Teacher Name"))

            # Determine Category (College / SHS)
            cat_val = clean_str(row.get("Category"))
            if "grade" in year_val.lower() or "shs" in cat_val.lower():
                cat_val = "SHS"
            else:
                cat_val = "College"

            # Determine exam_type and subject category
            inferred_exam_type, inferred_cat = classify_subject(sub_name)
            subj_cat = clean_str(row.get("Subject Category")).lower()
            if not subj_cat or subj_cat == "nan":
                subj_cat = inferred_cat
            elif inferred_cat == "general":
                subj_cat = "general"

            exam_type = clean_str(row.get("Exam Type")).lower()
            if not exam_type or exam_type == "nan":
                exam_type = inferred_exam_type
            elif inferred_exam_type == "practical":
                exam_type = "practical"

            row_data = [
                course_val,
                cat_val,
                year_val,
                section_val,
                code_val,
                sub_name,
                teacher_val,
                subj_cat,
                exam_type
            ]

            ws.row_dimensions[row_idx].height = 20
            is_alt = (row_idx % 2 == 1)

            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.border = thin_border
                if is_alt:
                    cell.fill = alt_fill
                
                # Alignments
                if col_idx in (2, 3, 4, 5, 8, 9):
                    cell.alignment = data_align_center
                else:
                    cell.alignment = data_align_left

            row_idx += 1

        # Adjust column widths
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save to both target locations
    wb.save(output_path_1)
    print(f"[OK] Saved aligned curriculum to: {output_path_1}")
    wb.save(output_path_2)
    print(f"[OK] Saved aligned curriculum to: {output_path_2}")

if __name__ == "__main__":
    generate_aligned_curriculum()
